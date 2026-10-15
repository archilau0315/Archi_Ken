"""Excel 报告渲染（M9）：openpyxl 生成「摘要 + 问题清单」双表，原子写入。"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SEVERITY_ORDER = ("error", "warning", "info")

I18N = {
    "zh": {
        "severity": {"error": "错误", "warning": "警告", "info": "提示"},
        "source": {"rule": "规则", "clash": "碰撞"},
        "title": "ArchiKen 模型检查报告",
        "model": "模型",
        "generated_at": "生成时间",
        "summary_sheet": "摘要",
        "issues_sheet": "问题清单",
        "summary": [
            "问题总数", "错误", "警告", "提示", "硬碰撞", "规则运行数", "未通过规则",
        ],
        "by_category": "按分类统计",
        "by_source": "按来源统计",
        "count": "数量",
        "headers": ["级别", "分类", "规则", "依据", "构件", "构件类型", "关联构件", "关联类型", "原因", "判定链", "位置", "来源"],
    },
    "en": {
        "severity": {"error": "Error", "warning": "Warning", "info": "Info"},
        "source": {"rule": "Rule", "clash": "Clash"},
        "title": "ArchiKen Model Check Report",
        "model": "Model",
        "generated_at": "Generated at",
        "summary_sheet": "Summary",
        "issues_sheet": "Issues",
        "summary": [
            "Total Issues", "Errors", "Warnings", "Infos", "Hard Clashes", "Rules Run", "Failed Rules",
        ],
        "by_category": "By Category",
        "by_source": "By Source",
        "count": "Count",
        "headers": ["Severity", "Category", "Rule", "Reference", "Element", "Element Type", "Related Element", "Related Type", "Reason", "Decision Chain", "Location", "Source"],
    },
}

HEAD_FILL = PatternFill("solid", fgColor="F2F4F7")
HEAD_FONT = Font(bold=True)
TITLE_FONT = Font(bold=True, size=14)
WRAP = Alignment(vertical="top", wrap_text=True)

WIDTHS = [8, 14, 22, 26, 26, 14, 26, 14, 42, 42, 24, 8]


def _fmt(value: Any) -> Any:
    """数值保留 3 位小数；None 转空串（Excel 单元格友好）。"""
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return round(value, 3)
    return value


def _location_text(location: Any) -> str:
    if isinstance(location, list) and location:
        return ", ".join(str(round(float(v), 3)) for v in location[:3])
    return ""


def _write_summary(book: Workbook, report: dict, title: str, lang: str = "zh") -> None:
    """摘要表：报告元信息 + 总体统计 + 按严重度 / 分类 / 来源统计。"""
    txt = I18N[lang if lang in I18N else "zh"]
    sheet = book.active
    sheet.title = txt["summary_sheet"]
    summary = report.get("summary") or {}

    sheet.append([title])
    sheet["A1"].font = TITLE_FONT
    sheet.append([txt["model"], report.get("model_path") or ""])
    sheet.append([txt["generated_at"], report.get("generated_at") or ""])
    sheet.append([])

    by_severity = summary.get("by_severity") or {}
    values = [
        summary.get("total", 0),
        by_severity.get("error", 0),
        by_severity.get("warning", 0),
        by_severity.get("info", 0),
        summary.get("clashes", 0),
        summary.get("rules_run", 0),
        summary.get("rules_failed", 0),
    ]
    for label, value in zip(txt["summary"], values):
        sheet.append([label, value])

    sheet.append([])
    sheet.append([txt["by_category"], txt["count"]])
    for category, count in sorted((summary.get("by_category") or {}).items(), key=lambda kv: -kv[1]):
        sheet.append([category, count])

    sheet.append([])
    sheet.append([txt["by_source"], txt["count"]])
    for source, count in sorted((summary.get("by_source") or {}).items()):
        sheet.append([txt["source"].get(source, source), count])

    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 60
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top")


def _write_issues(book: Workbook, issues: list[dict], lang: str = "zh") -> None:
    """问题清单表：表头 + 逐问题行（级别按 错误→警告→提示 排序）。"""
    txt = I18N[lang if lang in I18N else "zh"]
    headers = txt["headers"]
    sheet = book.create_sheet(txt["issues_sheet"])
    sheet.append(headers)
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        sheet.column_dimensions[get_column_letter(index)].width = WIDTHS[index - 1]

    ordered = sorted(issues, key=lambda issue: SEVERITY_ORDER.index(issue.get("severity") or "info")
                     if (issue.get("severity") or "info") in SEVERITY_ORDER else 99)
    for issue in ordered:
        element = issue.get("element") or {}
        related = issue.get("related") or {}
        sheet.append(
            [
                txt["severity"].get(issue.get("severity") or "", issue.get("severity") or ""),
                issue.get("category") or "",
                issue.get("title") or "",
                issue.get("reference") or "",
                element.get("name") or element.get("guid") or "",
                element.get("ifc_type") or "",
                related.get("name") or related.get("guid") or "",
                related.get("ifc_type") or "",
                issue.get("reason") or "",
                "\n".join(str(step) for step in issue.get("steps") or []),
                _location_text(issue.get("location")),
                txt["source"].get(issue.get("source") or "", issue.get("source") or ""),
            ]
        )
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=len(headers)):
        for cell in row:
            cell.alignment = WRAP
    sheet.freeze_panes = "A2"


def render_xlsx(report: dict, out: Path, title: str | None = None, lang: str = "zh") -> Path:
    """把统一问题清单渲染为 xlsx：摘要表 + 问题清单表。"""
    txt = I18N[lang if lang in I18N else "zh"]
    if title is None:
        title = txt["title"]
    book = Workbook()
    _write_summary(book, report, title, lang)
    _write_issues(book, report.get("issues") or [], lang)

    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_suffix(out.suffix + ".tmp")
    book.save(tmp)
    tmp.replace(out)
    return out
