"""信息提取 Excel 导出：openpyxl 写 xlsx，先写临时文件再原子替换。

支持分组模式：组标题行（合并 + 加粗底色）→ 组内明细 → 组小计行；末尾总合计行。
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.takeoff.columns import FIXED_FIELDS, FIXED_KEY_SET, ColumnError, localize_columns

I18N = {
    "zh": {
        "sheet": "信息提取",
        "total": "合计",
        "subtotal": "小计",
        "ungrouped": "(未分组)",
        "yes": "是",
        "no": "否",
    },
    "en": {
        "sheet": "Takeoff",
        "total": "Total",
        "subtotal": "Subtotal",
        "ungrouped": "(Ungrouped)",
        "yes": "Yes",
        "no": "No",
    },
}

_GROUP_FILL = PatternFill("solid", fgColor="E6EDFB")
_TOTAL_FILL = PatternFill("solid", fgColor="F2F4F7")
_BOLD = Font(bold=True)


def _column_header(column: dict) -> str:
    """列头带计量单位：面积 → 面积 (m²)；无单位则原样返回。"""
    unit = str(column.get("unit") or "").strip()
    return f"{column['name']} ({unit})" if unit else column["name"]


def export_xlsx(
    result: dict,
    out: Path,
    lang: str = "zh",
) -> Path:
    """把 run_takeoff 结果写为 xlsx：表头 + 数据行（可选分组） + 合计行。"""
    txt = I18N[lang if lang in I18N else "zh"]
    columns: list[dict] = result.get("columns") or []
    # 列名已按语言本地化（由路由层传入 localize_columns 结果）。
    headers = [field["name"] for field in localize_columns(FIXED_FIELDS, lang)] + [
        _column_header(column) for column in columns
    ]
    keys = [field["key"] for field in FIXED_FIELDS] + [column["key"] for column in columns]
    if len(set(keys)) != len(FIXED_KEY_SET) + len(columns):
        raise ColumnError("存在与固定字段重名的自定义列")

    book = Workbook()
    sheet = book.active
    sheet.title = txt["sheet"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = _BOLD

    groups = result.get("groups")
    if groups:
        _write_groups(sheet, groups, keys, txt)
    else:
        for row in result.get("rows") or []:
            sheet.append(_row_cells(row, keys, txt))

    totals: dict = result.get("totals") or {}
    if totals:
        label = txt["total"] if not groups else f"══ {txt['total']} ══"
        _write_summary_row(sheet, label, totals, keys, txt, _TOTAL_FILL)

    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(len(str(header)) * 2 + 4, 10)
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_suffix(out.suffix + ".tmp")
    book.save(tmp)
    tmp.replace(out)
    return out


def _write_groups(sheet, groups: list[dict], keys: list[str], txt: dict) -> None:
    """按组写：组标题行（跨列合并 + 加粗底色） → 组内明细 → 组小计行。"""
    span = len(keys)
    for group in groups:
        label = group.get("key") or txt["ungrouped"]
        header_idx = sheet.max_row + 1
        sheet.append([f"▼ {label}"])
        sheet.merge_cells(start_row=header_idx, start_column=1, end_row=header_idx, end_column=span)
        head_cell = sheet.cell(row=header_idx, column=1)
        head_cell.font = _BOLD
        head_cell.fill = _GROUP_FILL
        head_cell.alignment = Alignment(horizontal="left")
        for row in group.get("rows") or []:
            sheet.append(_row_cells(row, keys, txt))
        subtotals = group.get("subtotals") or {}
        _write_summary_row(sheet, f"  {txt['subtotal']} · {label}", subtotals, keys, txt, _TOTAL_FILL)


def _write_summary_row(sheet, label: str, values: dict, keys: list[str], txt: dict, fill) -> None:
    """合计 / 小计行：首列写标签，后续列写数值（非数值列留空），整行加粗 + 底色。"""
    row_idx = sheet.max_row + 1
    cells = [label] + [values.get(key, "") if isinstance(values.get(key), (int, float)) else "" for key in keys[1:]]
    sheet.append(cells)
    for col_idx in range(1, len(keys) + 1):
        cell = sheet.cell(row=row_idx, column=col_idx)
        cell.font = _BOLD
        cell.fill = fill


def _row_cells(row: dict, keys: list[str], txt: dict) -> list:
    """普通数据行：固定字段平铺取，自定义列从 values 嵌套字典取。"""
    values = row.get("values") or {}
    cells = []
    for key in keys:
        if key in row:
            cells.append(_cell(row.get(key), txt))
        else:
            cells.append(_cell(values.get(key), txt))
    return cells


def _cell(value: object, txt: dict) -> object:
    if isinstance(value, bool):
        return txt["yes"] if value else txt["no"]
    return value if value is not None else ""
